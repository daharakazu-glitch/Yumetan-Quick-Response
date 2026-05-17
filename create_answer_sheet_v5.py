"""
2026年度 論理・表現Ⅱ 中間テスト 解答用紙 v6
変更点:
  ・§3A: 1×10 縦長解答欄、前後の固定英文を印刷済み、生徒は並び替え部分のみ記入
  ・§3B: 各問の語数分のみ括弧()グループを1つのセルに表示
  ・A4×2ページ収録
"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MARGINS = 1.2
W  = 21 - 2 * MARGINS   # 18.6 cm
FS = 10.5

# ─────────────── ユーティリティ ───────────────

def set_row_h(row, h_cm, rule='exact'):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement('w:trHeight')
    e.set(qn('w:val'), str(int(h_cm * 567)))
    e.set(qn('w:hRule'), rule)
    trPr.append(e)

def font_run(run, size=FS, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = 'Arial'
    rpr = run._element.get_or_add_rPr()
    rf  = rpr.get_or_add_rFonts()
    rf.set(qn('w:ascii'),    'Arial')
    rf.set(qn('w:hAnsi'),    'Arial')
    rf.set(qn('w:eastAsia'), 'MS PGothic')

def cw(cell, text='', size=FS, bold=False,
       align=WD_ALIGN_PARAGRAPH.LEFT,
       valign=WD_ALIGN_VERTICAL.CENTER):
    cell.vertical_alignment = valign
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(0)
    r = p.add_run(text)
    font_run(r, size, bold)

def para(doc, text, size=FS, bold=False,
         align=WD_ALIGN_PARAGRAPH.LEFT,
         before=4, after=2):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after  = Pt(after)
    r = p.add_run(text)
    font_run(r, size, bold)

def no_b(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement('w:tcBorders')
    for s in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement(f'w:{s}')
        e.set(qn('w:val'), 'none'); e.set(qn('w:sz'), '0')
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), 'auto')
        b.append(e)
    tcPr.append(b)

def pgbrk(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(0)
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    p.add_run()._element.append(br)

def mktbl(doc, rows, cols, col_w, row_h=None, h_rule='exact'):
    t = doc.add_table(rows=rows, cols=cols)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    if isinstance(col_w, (int, float)):
        col_w = [col_w] * cols
    for j, w in enumerate(col_w):
        for i in range(rows):
            t.cell(i, j).width = Cm(w)
    if row_h is not None:
        heights = row_h if isinstance(row_h, list) else [row_h] * rows
        for i, h in enumerate(heights):
            set_row_h(t.rows[i], h, h_rule)
    return t

def set_right_tab(p_obj, pos_cm):
    """段落に右揃えタブを設定（セル左端からの距離）"""
    pPr = p_obj._p.get_or_add_pPr()
    tabs = OxmlElement('w:tabs')
    tab  = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'right')
    tab.set(qn('w:pos'), str(int(pos_cm * 567)))
    tabs.append(tab)
    pPr.append(tabs)

# ─────────────── §3A / §3B データ ───────────────

# §3A: [番号, 前文(固定), 後文(固定)]
# 前文・後文は解答用紙に印刷済み。生徒はその間の並び替えのみ記入。
S3A = [
    ['(1)',  '',                              'annually.'],
    ['(2)',  '',                              'in the world.'],
    ['(3)',  'It',                            'in summer.'],
    ['(4)',  'It is',                         'on the world map.'],
    ['(5)',  '',                              '.'],
    ['(6)',  'The author',                    'in her garden.'],
    ['(7)',  'It was quite frightening, but', '.'],
    ['(8)',  'In Japan,',                     '.'],
    ['(9)',  'Enjoy',                         '.'],
    ['(10)', 'In the magical town, she',      '.'],
]

# §3B: 各問の解答語数
B3B = [1, 2, 2, 3, 3, 3,   # Q1-6  (左列)
       1, 2, 2, 2, 2, 1]   # Q7-12 (右列)

def blank_parens(n):
    """n語分の括弧グループ（全角スペースで幅確保）"""
    if n == 1: return '(' + '　' * 16 + ')'
    if n == 2: return ('(' + '　' * 7 + ')') * 2
    if n == 3: return ('(' + '　' * 4 + ')') * 3
    return ''

# ─────────────── ドキュメント生成 ───────────────

def build(output_path):
    doc = Document()
    for sec in doc.sections:
        sec.page_width  = Cm(21)
        sec.page_height = Cm(29.7)
        for attr in ('left','right','top','bottom'):
            setattr(sec, f'{attr}_margin', Cm(MARGINS))

    # ── ヘッダー ──
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(3)
    r = p.add_run(
        '2026年度　山形県立山形南高等学校　２学年１学期中間テスト　論理・表現Ⅱ　解答用紙')
    font_run(r, 11, True)

    th = mktbl(doc, 2, 4, [2.0, 3.0, 1.8, 2.5])
    for i in range(2):
        set_row_h(th.rows[i], 0.72, 'atLeast')
    cw(th.cell(0,0), 'Class', 10, True, WD_ALIGN_PARAGRAPH.CENTER)
    cw(th.cell(0,2), 'No',    10, True, WD_ALIGN_PARAGRAPH.CENTER)
    cw(th.cell(1,0), 'Name',  10, True, WD_ALIGN_PARAGRAPH.CENTER)
    th.cell(1,1).merge(th.cell(1,3))

    # ══════════════ ページ１ ══════════════

    # ── §1 ──
    para(doc, '１　（思考・判断・表現　14点）', bold=True, before=6, after=2)

    t1ab = mktbl(doc, 1, 4, [3.2, 2.0, 3.2, 2.0], 0.84)
    cw(t1ab.cell(0,0), 'A（2点）', 9.5, True, WD_ALIGN_PARAGRAPH.CENTER)
    cw(t1ab.cell(0,2), 'B（2点）', 9.5, True, WD_ALIGN_PARAGRAPH.CENTER)

    para(doc, 'C（2×5=10点）', size=9.5, before=3, after=1)
    t1c = mktbl(doc, 1, 5, [3.7, 2.8, 1.9, 3.7, 3.7], 0.84)
    for j, q in enumerate(['(1)','(2)','(3)','(4)','(5)']):
        cw(t1c.cell(0,j), q, FS)

    # ── §2 ──
    para(doc, '２　（知識・技能　15点）', bold=True, before=5, after=2)

    para(doc, 'A（1×10=10点）　選択（A・B・C・D）', size=9.5, before=2, after=1)
    t2a = mktbl(doc, 2, 5, [3.0]*5, 0.84)
    for i in range(2):
        for j in range(5):
            cw(t2a.cell(i,j), f'({i*5+j+1})', FS)

    para(doc, 'B（1×5=5点）　英単語1語', size=9.5, before=3, after=1)
    t2b = mktbl(doc, 1, 5, [3.6]*5, 0.84)
    for j in range(5):
        cw(t2b.cell(0,j), f'({j+1})', FS)

    # ── §3 ──
    para(doc, '３　（知識・技能　28点）', bold=True, before=5, after=2)

    # §3A: 1列×10行、前後の固定英文を印刷
    para(doc, 'A（1×10=10点）　語句を並び替えて英文を完成させよ　完全正答のみ得点',
         size=9.5, before=2, after=1)

    NUM_W = 1.2          # 番号列幅 (cm)
    ANS_W = W - NUM_W    # 解答列幅 = 17.4 cm
    t3a = mktbl(doc, 10, 2, [NUM_W, ANS_W], 0.85)

    for i, (num, pre, post) in enumerate(S3A):
        # 番号セル
        cw(t3a.cell(i,0), num, FS, False, WD_ALIGN_PARAGRAPH.CENTER)
        # 解答セル：前文（左）＋タブ＋後文（右）
        cell = t3a.cell(i,1)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(0)
        # 右揃えタブをセル右端近くに設定
        set_right_tab(p, ANS_W - 0.15)
        r_pre = p.add_run((pre + '  ') if pre else '')
        font_run(r_pre)
        r_tab = p.add_run('\t')
        font_run(r_tab)
        r_post = p.add_run(post)
        font_run(r_post)

    # §3B: 各問の語数分だけ()グループ
    para(doc, 'B（1×12=12点）　空所に適切な語を入れよ', size=9.5, before=3, after=1)

    NW = 0.8                      # 番号列幅
    AW = (W - NW * 2) / 2        # 解答列幅 = (18.6-1.6)/2 = 8.5 cm
    t3b = mktbl(doc, 6, 4, [NW, AW, NW, AW], 0.93)

    for i in range(6):
        for side in range(2):
            qi  = i + side * 6
            nc  = side * 2
            ac  = side * 2 + 1
            cw(t3b.cell(i,nc), f'({qi+1})', 9.5, False, WD_ALIGN_PARAGRAPH.CENTER)
            cw(t3b.cell(i,ac), blank_parens(B3B[qi]), FS, False,
               WD_ALIGN_PARAGRAPH.CENTER)

    # §3C
    para(doc, 'C（1×6=6点）　動詞を正しい形に変えよ', size=9.5, before=3, after=1)
    t3c = mktbl(doc, 1, 6, [3.0]*6, 0.84)
    for j in range(6):
        cw(t3c.cell(0,j), f'({j+1})', FS)

    # ══════════════ ページ区切り ══════════════
    pgbrk(doc)

    # ── §4 ──
    para(doc, '４　（思考・判断・表現　12点）', bold=True, before=0, after=2)
    t4 = mktbl(doc, 6, 1, [W], 0.90)
    for i in range(6):
        cw(t4.cell(i,0), f'({i+1})', FS)

    # ── §5 ──
    para(doc, '５　（思考・判断・表現　10点）', bold=True, before=4, after=2)
    para(doc, 'A（1×5=5点）　　　　　　　　　　　　　B（1×5=5点）',
         size=9.5, before=1, after=1)
    HW5 = W / 2
    t5 = mktbl(doc, 5, 2, [HW5, HW5], 0.84)
    for i in range(5):
        for j in range(2):
            cw(t5.cell(i,j), f'({i+1})', FS)

    # ── §6 ──
    para(doc, '６　（思考・判断・表現　15点）', bold=True, before=3, after=2)

    para(doc, '問1（1点）・問2（2点）・問5（2点）　各 A〜D から選択',
         size=9.5, before=2, after=1)
    t6s = mktbl(doc, 1, 3, [3.2, 3.2, 3.2], 0.84)
    for j, lbl in enumerate(['問1（1点）','問2（2点）','問5（2点）']):
        cw(t6s.cell(0,j), lbl, 9.5, True, WD_ALIGN_PARAGRAPH.CENTER)

    para(doc, '問3（1×5=5点）　S＝スパルタ　A＝アテネ', size=9.5, before=3, after=1)
    t6q3 = mktbl(doc, 1, 5, [2.4]*5, 0.84)
    for j, lbl in enumerate(['(i)','(ii)','(iii)','(iv)','(v)']):
        cw(t6q3.cell(0,j), lbl, FS, False, WD_ALIGN_PARAGRAPH.CENTER)

    para(doc, '問4（2点）　スパルタの女性がアテネの女性より「自由」を得られた理由',
         size=9.5, before=3, after=1)
    t6q4 = mktbl(doc, 2, 1, [W], 0.88)
    for i in range(2): cw(t6q4.cell(i,0), '')

    para(doc, '問6（3点）　下線部(2)の日本語訳', size=9.5, before=3, after=1)
    t6q6 = mktbl(doc, 2, 1, [W], 0.88)
    for i in range(2): cw(t6q6.cell(i,0), '')

    # ── §7 ──
    para(doc, '７　（思考・判断・表現　6点）　50語以上', bold=True, before=4, after=2)
    t7 = mktbl(doc, 6, 1, [W], 0.83)
    for i in range(6): cw(t7.cell(i,0), '')

    # ── 得点集計 ──
    para(doc, '得点集計', bold=True, before=4, after=1)
    ts = mktbl(doc, 2, 3, [5.8, 7.2, 5.6])
    set_row_h(ts.rows[0], 0.60)
    set_row_h(ts.rows[1], 1.20)
    for j, lbl in enumerate(['２・３（知識・技能）　/43',
                              '１・４〜７（思考・判断・表現）　/57',
                              '合計　/100']):
        cw(ts.cell(0,j), lbl, 9, True, WD_ALIGN_PARAGRAPH.CENTER)
        cw(ts.cell(1,j), '',  FS,      align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.save(output_path)
    print(f'Saved: {output_path}')


if __name__ == '__main__':
    build('/home/user/Yumetan-Quick-Response/2026年度_論理表現Ⅱ_中間テスト_解答用紙.docx')
